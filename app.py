import io
import re
import os

import openpyxl
import pandas as pd
from flask import Flask, jsonify, render_template, request

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 16 * 1024 * 1024  # 16 MB


# ── Generic bid column detection ─────────────────────────────────────────────
ALIASES = {
    "Bid_ID":        ["bid_id", "bid id", "bidid", "id", "ref", "reference",
                      "quote_id", "quoteid", "rfq", "tender_id"],
    "Supplier":      ["supplier", "vendor", "company", "bidder", "contractor",
                      "supplier_name", "vendor_name", "firm"],
    "Category":      ["category", "cat", "type", "group", "division",
                      "department", "commodity"],
    "Product":       ["product", "item", "description", "material", "goods",
                      "service", "item_description", "product_name"],
    "Quantity":      ["quantity", "qty", "units", "volume", "count",
                      "amount", "no_of_units", "num"],
    "Unit_Price":    ["unit_price", "unit price", "price", "rate", "cost",
                      "unit_cost", "per_unit", "unitprice", "unit rate"],
    "Total_Amount":  ["total_amount", "total amount", "total", "total_cost",
                      "grand_total", "subtotal", "line_total", "ext_price"],
    "Delivery_Days": ["delivery_days", "delivery days", "delivery",
                      "lead_time", "leadtime", "days", "turnaround",
                      "ship_days", "eta_days"],
    "Proposal_Type": ["proposal_type", "proposal type", "bid_type",
                      "proposal", "type", "quote_type"],
}


def _norm(s: str) -> str:
    return re.sub(r"[\s_\-]+", "_", s.strip().lower())


def detect_columns(df: pd.DataFrame) -> dict:
    df_cols = {_norm(c): c for c in df.columns}
    mapping = {}
    for field, aliases in ALIASES.items():
        for alias in aliases:
            if _norm(alias) in df_cols:
                mapping[field] = df_cols[_norm(alias)]
                break
    return mapping


def clean_number(val) -> float:
    try:
        return float(re.sub(r"[£$,\s]", "", str(val)))
    except (ValueError, TypeError):
        return 0.0


def parse_df(df: pd.DataFrame, col_map: dict) -> list:
    records = []
    for i, row in df.iterrows():
        def get(field, default=""):
            col = col_map.get(field)
            return row[col] if col else default

        unit_price = clean_number(get("Unit_Price", 0))
        quantity   = clean_number(get("Quantity", 1)) or 1
        total      = clean_number(get("Total_Amount", 0))
        if total == 0 and unit_price > 0:
            total = unit_price * quantity

        records.append({
            "Bid_ID":        str(get("Bid_ID", f"BID{i+1:03d}")) or f"BID{i+1:03d}",
            "Supplier":      str(get("Supplier", "Unknown")),
            "Category":      str(get("Category", "General")),
            "Product":       str(get("Product", f"Item {i+1}")),
            "Quantity":      quantity,
            "Unit_Price":    unit_price,
            "Total_Amount":  total,
            "Delivery_Days": int(clean_number(get("Delivery_Days", 0))),
            "Proposal_Type": str(get("Proposal_Type", "Original")).strip() or "Original",
        })
    return [r for r in records if r["Supplier"] not in ("Unknown", "", "nan")]


def read_csv_content(content: bytes) -> pd.DataFrame:
    text = content.decode("utf-8-sig", errors="replace")
    return pd.read_csv(io.StringIO(text))


def read_any_table_file(content: bytes, filename: str) -> pd.DataFrame:
    ext = os.path.splitext(filename.lower())[1]
    if ext in {".xls", ".xlsx", ".xlsm", ".xlsb", ".ods"}:
        return pd.read_excel(io.BytesIO(content))
    if ext == ".csv" or not ext:
        return read_csv_content(content)
    try:
        return read_csv_content(content)
    except Exception as csv_exc:
        try:
            return pd.read_excel(io.BytesIO(content))
        except Exception as excel_exc:
            raise ValueError(
                f"Could not parse as CSV or Excel. CSV: {csv_exc}; Excel: {excel_exc}"
            ) from excel_exc


# ── QCS (Quote Comparison Sheet) parser ──────────────────────────────────────

def _clean_price(v):
    """Return float or None from a cell value; ignores zeros."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v) if v != 0 else None
    s = re.sub(r"[,$\s]", "", str(v))
    try:
        f = float(s)
        return f if f != 0 else None
    except (ValueError, TypeError):
        return None


def _is_qcs(content: bytes) -> bool:
    """Return True if the Excel file looks like a QCS document."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        hits = 0
        for row in ws.iter_rows(max_row=35, values_only=True):
            for cell in row:
                if cell:
                    s = str(cell)
                    if "QUOTE COMPARISON SHEET" in s or "TOTAL PROPOSAL PRICE" in s:
                        hits += 1
            if hits >= 1:
                return True
        return False
    except Exception:
        return False


def parse_qcs_excel(content: bytes) -> dict:
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active

    # Read everything as a flat list of rows
    rows = [list(r) for r in ws.iter_rows(values_only=True)]

    # ── Locate key rows ───────────────────────────────────────────────────
    def find_row(text, max_r=80):
        for i, row in enumerate(rows[:max_r]):
            if any(cell is not None and text in str(cell) for cell in row):
                return i
        return None

    r_bidder  = find_row("BIDDER 1")
    r_legal   = find_row("Bidder Legal Name")
    r_biddate = find_row("Bid Date or")
    r_valid   = find_row("Bid Valid Until Date")
    r_valdays = find_row("Bid Validity Duration")
    r_type    = find_row("Firm or Budgetary")
    r_total   = find_row("TOTAL PROPOSAL PRICE")
    r_pb      = find_row("Pricing Breakdown")
    r_teval   = find_row("TOTAL EVALUATED PRICE")

    # ── Find bidder column positions from "BIDDER N" labels ───────────────
    bidder_start = {}
    if r_bidder is not None:
        for j, cell in enumerate(rows[r_bidder]):
            if cell:
                m = re.search(r"BIDDER\s*(\d+)", str(cell))
                if m:
                    n = int(m.group(1))
                    if n not in bidder_start:
                        bidder_start[n] = j

    num_bidders = len(bidder_start)

    # ── Find exact unit-price / total-price column positions ─────────────
    up_cols = []
    tp_cols = []

    if r_pb is not None:
        for offset in range(1, 5):
            idx = r_pb + offset
            if idx >= len(rows):
                break
            row = rows[idx]
            found_up, found_tp = [], []
            for j, cell in enumerate(row):
                if cell:
                    s = str(cell)
                    if "Unit Price" in s:
                        found_up.append(j)
                    if "Total Price" in s:
                        found_tp.append(j)
            if found_up and found_tp:
                up_cols = sorted(found_up)
                tp_cols = sorted(found_tp)
                break

    # Fallback: derive from BIDDER N start cols (UP=start, TOTAL=start+2)
    if not up_cols and bidder_start:
        for n in sorted(bidder_start):
            up_cols.append(bidder_start[n])
            tp_cols.append(bidder_start[n] + 2)

    # ── Helper: get one value per bidder from a row ───────────────────────
    def per_bidder(row_idx, cols, convert=None):
        result = {}
        if row_idx is None or row_idx >= len(rows):
            return result
        row = rows[row_idx]
        for i, col in enumerate(cols):
            val = None
            for off in range(-1, 5):
                c = col + off
                if 0 <= c < len(row) and row[c] is not None:
                    s = str(row[c]).strip()
                    if s and s not in ("None", "-", ""):
                        val = row[c]
                        break
            if convert and val is not None:
                try:
                    val = convert(val)
                except Exception:
                    val = None
            result[i + 1] = val
        return result

    def _str(v):
        if v is None:
            return None
        s = str(v).strip()
        return s if s not in ("None", "-", "") else None

    def _int(v):
        if v is None:
            return None
        try:
            return int(float(str(v)))
        except Exception:
            return None

    # ── Project metadata ──────────────────────────────────────────────────
    proj = {}
    for i, row in enumerate(rows[:35]):
        for j, cell in enumerate(row):
            if not cell:
                continue
            s = str(cell)
            if "Project Name" in s and "Bidder" not in s:
                nxt = rows[i][j + 1] if j + 1 < len(rows[i]) else None
                if nxt:
                    proj.setdefault("name", str(nxt).strip())
            if "Job Site" in s and "Bidder" not in s:
                nxt = rows[i][j + 1] if j + 1 < len(rows[i]) else None
                if nxt:
                    proj.setdefault("site", str(nxt).strip())
            if "Engineering" in s and ":" not in s:
                for k in range(j + 1, min(j + 4, len(rows[i]))):
                    v = rows[i][k]
                    if v and "Engineering" not in str(v):
                        proj.setdefault("engineering", str(v).strip())
                        break
            if "BID DUE" in s:
                for k in range(j + 1, min(j + 8, len(rows[i]))):
                    v = rows[i][k]
                    if v and str(v).strip() not in ("None", "-", ""):
                        proj.setdefault("bid_due", str(v).strip())
                        break

    # ── Bidder names ──────────────────────────────────────────────────────
    names = {}
    if r_legal is not None and up_cols:
        row = rows[r_legal]
        for i, col in enumerate(up_cols):
            for off in range(-1, 6):
                c = col + off
                if 0 <= c < len(row) and row[c] is not None:
                    s = str(row[c]).strip()
                    if (s and s not in ("None", "-", "")
                            and not s.startswith("$")
                            and not re.match(r"^[\d,\.]+$", s)
                            and "Bidder Legal" not in s):
                        names[i + 1] = s
                        break

    # ── Per-bidder commercial data ────────────────────────────────────────
    bid_dates  = per_bidder(r_biddate, up_cols, _str)
    bid_valid  = per_bidder(r_valid,   up_cols, _str)
    valid_days = per_bidder(r_valdays, up_cols, _int)
    bid_types  = per_bidder(r_type,    up_cols, _str)

    # Total proposal prices — search within each bidder's own 3-col block only
    totals_map = {}
    if r_total is not None:
        row = rows[r_total]
        for i, col in enumerate(up_cols):
            # Search UP col + next 2 cols (definition + total), then one before
            search = [col, col + 1, col + 2, col - 1]
            for c in search:
                if 0 <= c < len(row):
                    p = _clean_price(row[c])
                    if p and p > 100_000:
                        totals_map[i + 1] = p
                        break

    # ── Line items ────────────────────────────────────────────────────────
    UNITS    = {"LF", "LOT", "DAY", "RT", "EA", "FT", "LB", "KG", "TON", "EACH"}
    SKIP_KWS = {
        "Pricing for", "Commercial Adjustments", "Technical Adjustments",
        "TOTAL UNIT PRICE", "Security Cost Per", "Kiewit Evaluation",
        "Pricing Reliability", "Total Proposal Price Comparison",
        "Amount above minimum", "Percent above", "Awarding",
        "Technically Acceptable", "Commercially Acceptable",
        "Price Carried", "Alternate Price", "IMPORTANT:", "Options Excluded",
        "Financial Security", "Warranty", "Shop Details", "Shipping",
        "Lead Time", "Bidder Certification", "Insurance and Quality",
        "Proposal Exception", "Progress Payment",
    }

    line_items = []

    if r_pb is not None and up_cols:
        section    = "General"
        data_start = r_pb + 2
        data_end   = r_teval if r_teval else len(rows)

        for i in range(data_start, data_end):
            row = rows[i]
            if not any(c is not None for c in row):
                continue

            row_text = " ".join(str(c) for c in row if c is not None)

            # Section header?
            if "Pricing for" in row_text:
                for cell in row:
                    if cell and "Pricing for" in str(cell):
                        section = str(cell).strip()
                        break
                continue

            # Non-data rows
            if any(kw in row_text for kw in SKIP_KWS):
                continue

            # Extract description and unit/qty from left columns
            desc = None
            unit = None
            qty  = None

            for j in range(min(7, len(row))):
                v = row[j]
                if v is None:
                    continue
                s = str(v).strip()
                if not s or s in ("None", "-"):
                    continue
                if s in UNITS:
                    unit = s
                    continue
                if re.match(r"^\d{1,4}$", s):            # item # (1-4 digits)
                    continue
                if re.match(r"^\d[\d,]+$", s):            # big number = qty
                    try:
                        qty = float(s.replace(",", ""))
                    except Exception:
                        pass
                    continue
                if re.match(r"^\d+\.\d+\.\d+", s):        # spec code
                    continue
                if s.startswith("$"):
                    continue
                if len(s) > 3 and desc is None:
                    desc = s

            if not desc:
                continue

            # Scan slightly wider for unit / qty if still missing
            if not unit:
                for j in range(3, min(11, len(row))):
                    v = row[j]
                    if v and str(v).strip() in UNITS:
                        unit = str(v).strip()
                        break

            if qty is None:
                for j in range(2, min(10, len(row))):
                    v = row[j]
                    if isinstance(v, (int, float)) and v > 0:
                        qty = float(v)
                        break

            # Per-bidder prices
            prices = {}
            totals = {}
            defs   = {}

            for bi, (uc, tc) in enumerate(zip(up_cols, tp_cols)):
                bn = str(bi + 1)
                up = row[uc] if uc < len(row) else None
                tp = row[tc] if tc < len(row) else None
                # Definition column (between UP and TP)
                dc = uc + 1
                if dc != tc and dc < len(row):
                    defs[bn] = str(row[dc]).strip() if row[dc] else None
                prices[bn] = _clean_price(up)
                totals[bn] = _clean_price(tp)

            if not any(v is not None for v in prices.values()):
                continue

            line_items.append({
                "description": desc,
                "section":     section,
                "unit":        unit,
                "qty":         qty,
                "prices":      prices,
                "totals":      totals,
                "definitions": defs,
            })

    # ── Assemble result ───────────────────────────────────────────────────
    bidders = []
    for n in range(1, num_bidders + 1):
        bidders.append({
            "num":             n,
            "name":            names.get(n) or f"Bidder {n}",
            "bid_date":        bid_dates.get(n),
            "bid_valid_until": bid_valid.get(n),
            "validity_days":   valid_days.get(n),
            "bid_type":        bid_types.get(n),
            "total_price":     totals_map.get(n, 0),
        })

    return {
        "type":        "qcs",
        "project":     proj,
        "bidders":     bidders,
        "line_items":  line_items,
        "num_bidders": num_bidders,
    }


# ── Routes ────────────────────────────────────────────────────────────────────
@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/upload", methods=["POST"])
def upload():
    if "file" not in request.files:
        return jsonify({"error": "No file provided"}), 400

    f = request.files["file"]
    content = f.read()
    filename = f.filename or ""
    ext = os.path.splitext(filename.lower())[1]

    # QCS detection for Excel files
    if ext in {".xlsx", ".xlsm", ".xlsb", ".xls", ".ods"}:
        if _is_qcs(content):
            try:
                result = parse_qcs_excel(content)
                return jsonify(result)
            except Exception as exc:
                return jsonify({"error": f"QCS parse failed: {exc}"}), 400

    # Generic parse path
    try:
        df = read_any_table_file(content, filename)
    except Exception as exc:
        return jsonify({"error": f"Could not parse file: {exc}"}), 400

    df.dropna(how="all", inplace=True)
    col_map = detect_columns(df)

    if "Supplier" not in col_map and "Product" not in col_map:
        return jsonify({
            "error":   "no_match",
            "columns": list(df.columns),
            "preview": df.head(3).fillna("").to_dict(orient="records"),
        }), 422

    records  = parse_df(df, col_map)
    detected = dict(col_map)

    return jsonify({
        "type":             "generic",
        "records":          records,
        "columns_detected": detected,
        "total_rows":       len(records),
        "unmatched_cols":   [c for c in df.columns if c not in col_map.values()],
    })


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=False)
